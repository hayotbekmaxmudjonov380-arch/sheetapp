import sqlite3
from flask import Flask, render_template_string, request, redirect, url_for, send_file
import openpyxl
import io
import os
from datetime import datetime

app = Flask(__name__)
DB_PATH = "restaurant.db"

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

# Updated HTML Template with Horizontal Top Navbar
HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="uz">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Restoran Hisob Dasturi</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
    <style>
        body { background-color: #f8f9fa; font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; }
        .navbar-brand { font-weight: bold; font-size: 1.25rem; }
        .card-stat { border-left: 4px solid #0d6efd; }
    </style>
</head>
<body>

<!-- Top Navigation Bar -->
<nav class="navbar navbar-expand-lg navbar-dark bg-dark shadow-sm px-4">
    <a class="navbar-brand" href="/">🍽️ Restoran Hisob</a>
    <button class="navbar-toggler" type="button" data-bs-toggle="collapse" data-bs-target="#navbarNav">
        <span class="navbar-toggler-icon"></span>
    </button>
    <div class="collapse navbar-collapse justify-content-between" id="navbarNav">
        <ul class="navbar-nav">
            <li class="nav-item">
                <a class="nav-link {% if active_tab == 'dashboard' %}active text-white fw-bold bg-primary rounded px-3{% endif %}" href="/">📊 Bosh sahifa</a>
            </li>
            <li class="nav-item mx-1">
                <a class="nav-link {% if active_tab == 'employees' %}active text-white fw-bold bg-primary rounded px-3{% endif %}" href="/employees">👥 Xodimlar</a>
            </li>
            <li class="nav-item mx-1">
                <a class="nav-link {% if active_tab == 'salary' %}active text-white fw-bold bg-primary rounded px-3{% endif %}" href="/salary">💰 Ish haqi</a>
            </li>
            <li class="nav-item mx-1">
                <a class="nav-link {% if active_tab == 'transactions' %}active text-white fw-bold bg-primary rounded px-3{% endif %}" href="/transactions">📥 Kirim / Chiqim</a>
            </li>
            <li class="nav-item mx-1">
                <a class="nav-link {% if active_tab == 'balance' %}active text-white fw-bold bg-primary rounded px-3{% endif %}" href="/balance">📈 Balans va DDS</a>
            </li>
        </ul>
        <div>
            <a href="/export" class="btn btn-success btn-sm text-white fw-bold">📥 Excel'ga Yuklab olish</a>
        </div>
    </div>
</nav>

<!-- Main Content Container -->
<div class="container my-4">
    {% with messages = get_flashed_messages(with_categories=true) %}
      {% if messages %}
        {% for category, message in messages %}
          <div class="alert alert-{{ 'success' if category == 'success' else 'danger' }} alert-dismissible fade show" role="alert">
            {{ message }}
            <button type="button" class="btn-close" data-bs-dismiss="alert" aria-label="Close"></button>
          </div>
        {% endfor %}
      {% endif %}
    {% endwith %}

    {% if active_tab == 'dashboard' %}
        <h2 class="mb-4">Boshqaruv Paneli (Dashboard)</h2>
        <div class="row">
            <div class="col-md-3 mb-3">
                <div class="card p-3 shadow-sm card-stat">
                    <h5>Xodimlar</h5>
                    <h3>{{ emp_count }} ta</h3>
                </div>
            </div>
            <div class="col-md-3 mb-3">
                <div class="card p-3 shadow-sm card-stat border-success">
                    <h5>Jami Kirim</h5>
                    <h3 class="text-success">{{ "{:,.0f}".format(total_kirim) }} so'm</h3>
                </div>
            </div>
            <div class="col-md-3 mb-3">
                <div class="card p-3 shadow-sm card-stat border-danger">
                    <h5>Jami Chiqim</h5>
                    <h3 class="text-danger">{{ "{:,.0f}".format(total_chiqim) }} so'm</h3>
                </div>
            </div>
            <div class="col-md-3 mb-3">
                <div class="card p-3 shadow-sm card-stat border-primary">
                    <h5>Sof Foyda</h5>
                    <h3 class="text-primary">{{ "{:,.0f}".format(sof_foyda) }} so'm</h3>
                </div>
            </div>
        </div>

    {% elif active_tab == 'employees' %}
        <div class="d-flex justify-content-between align-items-center mb-3">
            <h2>Xodimlar ro'yxati</h2>
            <button class="btn btn-primary" data-bs-toggle="modal" data-bs-target="#addEmpModal">➕ Yangi xodim qo'shish</button>
        </div>
        
        <div class="card shadow-sm p-3">
            <table class="table table-striped table-hover">
                <thead>
                    <tr>
                        <th>ID</th>
                        <th>Ism familiya</th>
                        <th>Lavozimi</th>
                        <th>Telefon</th>
                        <th>Tug'ilgan kun</th>
                        <th>Status</th>
                    </tr>
                </thead>
                <tbody>
                    {% for emp in employees %}
                    <tr>
                        <td>{{ emp.id }}</td>
                        <td>{{ emp.name }}</td>
                        <td>{{ emp.position }}</td>
                        <td>{{ emp.phone }}</td>
                        <td>{{ emp.birth_date }}</td>
                        <td><span class="badge bg-success">{{ emp.status }}</span></td>
                    </tr>
                    {% endfor %}
                </tbody>
            </table>
        </div>

        <!-- Add Employee Modal -->
        <div class="modal fade" id="addEmpModal" tabindex="-1">
            <div class="modal-dialog">
                <div class="modal-content">
                    <form method="POST" action="/employees/add">
                        <div class="modal-header">
                            <h5 class="modal-title">Yangi xodim qo'shish</h5>
                            <button type="button" class="btn-close" data-bs-dismiss="modal"></button>
                        </div>
                        <div class="modal-body">
                            <div class="mb-3">
                                <label class="form-label">Ism familiya</label>
                                <input type="text" class="form-control" name="name" required>
                            </div>
                            <div class="mb-3">
                                <label class="form-label">Lavozimi</label>
                                <input type="text" class="form-control" name="position">
                            </div>
                            <div class="mb-3">
                                <label class="form-label">Telefon</label>
                                <input type="text" class="form-control" name="phone">
                            </div>
                        </div>
                        <div class="modal-footer">
                            <button type="submit" class="btn btn-primary">Saqlash</button>
                        </div>
                    </form>
                </div>
            </div>
        </div>

    {% elif active_tab == 'salary' %}
        <h2 class="mb-3">Kunlik Ish Haqi Yozuvlari</h2>
        <div class="card shadow-sm p-3">
            <table class="table table-striped table-hover">
                <thead>
                    <tr>
                        <th>ID</th>
                        <th>Xodim</th>
                        <th>Sana</th>
                        <th>Kirish</th>
                        <th>Chiqim</th>
                        <th>Kunbay summa</th>
                        <th>Ishlagan vaqti</th>
                    </tr>
                </thead>
                <tbody>
                    {% for sal in salaries %}
                    <tr>
                        <td>{{ sal.id }}</td>
                        <td>{{ sal.employee_name }}</td>
                        <td>{{ sal.work_date }}</td>
                        <td>{{ sal.check_in }}</td>
                        <td>{{ sal.check_out }}</td>
                        <td><strong>{{ "{:,.0f}".format(sal.amount) }} so'm</strong></td>
                        <td>{{ sal.worked_hours }}</td>
                    </tr>
                    {% endfor %}
                </tbody>
            </table>
        </div>

    {% elif active_tab == 'transactions' %}
        <div class="d-flex justify-content-between align-items-center mb-3">
            <h2>Kirim va Chiqimlar</h2>
            <button class="btn btn-primary" data-bs-toggle="modal" data-bs-target="#addTxModal">➕ Kirim / Chiqim qo'shish</button>
        </div>
        
        <div class="card shadow-sm p-3">
            <table class="table table-striped table-hover">
                <thead>
                    <tr>
                        <th>ID</th>
                        <th>Turi</th>
                        <th>Modda / Kategoriya</th>
                        <th>Kontragent</th>
                        <th>Summa</th>
                        <th>Sana</th>
                    </tr>
                </thead>
                <tbody>
                    {% for tx in transactions %}
                    <tr>
                        <td>{{ tx.id }}</td>
                        <td>
                            {% if tx.type == 'kirim' %}
                                <span class="badge bg-success">Kirim 🟢</span>
                            {% else %}
                                <span class="badge bg-danger">Chiqim 🔴</span>
                            {% endif %}
                        </td>
                        <td>{{ tx.category }}</td>
                        <td>{{ tx.kontragent }}</td>
                        <td><strong>{{ "{:,.0f}".format(tx.amount) }} so'm</strong></td>
                        <td>{{ tx.tx_date }}</td>
                    </tr>
                    {% endfor %}
                </tbody>
            </table>
        </div>

        <!-- Add Transaction Modal -->
        <div class="modal fade" id="addTxModal" tabindex="-1">
            <div class="modal-dialog">
                <div class="modal-content">
                    <form method="POST" action="/transactions/add">
                        <div class="modal-header">
                            <h5 class="modal-title">Yangi Kirim / Chiqim qo'shish</h5>
                            <button type="button" class="btn-close" data-bs-dismiss="modal"></button>
                        </div>
                        <div class="modal-body">
                            <div class="mb-3">
                                <label class="form-label">Turi</label>
                                <select class="form-select" name="type">
                                    <option value="kirim">Kirim</option>
                                    <option value="chiqim">Chiqim</option>
                                </select>
                            </div>
                            <div class="mb-3">
                                <label class="form-label">Modda / Kategoriya (masalan: Savdo, Tovar, Arenda...)</label>
                                <input type="text" class="form-control" name="category" required>
                            </div>
                            <div class="mb-3">
                                <label class="form-label">Kontragent</label>
                                <input type="text" class="form-control" name="kontragent">
                            </div>
                            <div class="mb-3">
                                <label class="form-label">Summa (so'm)</label>
                                <input type="number" step="any" class="form-control" name="amount" required>
                            </div>
                            <div class="mb-3">
                                <label class="form-label">Sana</label>
                                <input type="date" class="form-control" name="tx_date" value="{{ today }}" required>
                            </div>
                        </div>
                        <div class="modal-footer">
                            <button type="submit" class="btn btn-primary">Saqlash</button>
                        </div>
                    </form>
                </div>
            </div>
        </div>

    {% elif active_tab == 'balance' %}
        <h2 class="mb-3">Balans va Sof Foyda Hisoboti</h2>
        <div class="card shadow-sm p-4 col-md-8">
            <table class="table">
                <tr>
                    <td><strong>Savdo:</strong></td>
                    <td class="text-end text-success"><strong>{{ "{:,.0f}".format(balance.savdo) }} so'm</strong></td>
                </tr>
                <tr>
                    <td><strong>Sotib olingan tovar:</strong></td>
                    <td class="text-end text-danger"><strong>{{ "{:,.0f}".format(balance.tovar) }} so'm</strong></td>
                </tr>
                <tr class="table-secondary">
                    <td><strong>Yalpi foyda (Savdo - Tovar):</strong></td>
                    <td class="text-end"><strong>{{ "{:,.0f}".format(balance.yalpi_foyda) }} so'm</strong></td>
                </tr>
                <tr>
                    <td>Ish haqi:</td>
                    <td class="text-end">{{ "{:,.0f}".format(balance.ish_haqi) }} so'm</td>
                </tr>
                <tr>
                    <td>Kommunal:</td>
                    <td class="text-end">{{ "{:,.0f}".format(balance.kommunal) }} so'm</td>
                </tr>
                <tr>
                    <td>Arenda:</td>
                    <td class="text-end">{{ "{:,.0f}".format(balance.arenda) }} so'm</td>
                </tr>
                <tr>
                    <td>Marketing:</td>
                    <td class="text-end">{{ "{:,.0f}".format(balance.marketing) }} so'm</td>
                </tr>
                <tr>
                    <td>Remont:</td>
                    <td class="text-end">{{ "{:,.0f}".format(balance.remont) }} so'm</td>
                </tr>
                <tr>
                    <td>Nalog (Savdoning 2%):</td>
                    <td class="text-end">{{ "{:,.0f}".format(balance.nalog) }} so'm</td>
                </tr>
                <tr class="table-warning">
                    <td><strong>Jami xarajat:</strong></td>
                    <td class="text-end"><strong>{{ "{:,.0f}".format(balance.jami_xarajat) }} so'm</strong></td>
                </tr>
                <tr class="table-primary">
                    <td><h4>Sof foyda:</h4></td>
                    <td class="text-end"><h4>{{ "{:,.0f}".format(balance.sof_foyda) }} so'm</h4></td>
                </tr>
            </table>
        </div>
    {% endif %}
</div>
<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
</body>
</html>
"""

@app.route('/')
def dashboard():
    conn = get_db()
    emp_count = conn.execute("SELECT COUNT(*) FROM employees").fetchone()[0]
    total_kirim = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='kirim'").fetchone()[0] or 0
    total_chiqim = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim'").fetchone()[0] or 0
    
    savdo = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='kirim' AND (category LIKE '%savdo%' OR category LIKE '%Савдо%')").fetchone()[0] or 0
    tovar = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND (category LIKE '%tovar%' OR category LIKE '%товар%')").fetchone()[0] or 0
    ish_haqi = conn.execute("SELECT SUM(amount) FROM work_records").fetchone()[0] or 0
    kommunal = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND category LIKE '%kommunal%'").fetchone()[0] or 0
    arenda = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND category LIKE '%arenda%'").fetchone()[0] or 0
    marketing = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND category LIKE '%marketing%'").fetchone()[0] or 0
    remont = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND category LIKE '%remont%'").fetchone()[0] or 0
    nalog = savdo * 0.02
    yalpi = savdo - tovar
    jami_xarajat = ish_haqi + kommunal + arenda + marketing + remont + nalog
    sof_foyda = yalpi - jami_xarajat
    
    conn.close()
    return render_template_string(HTML_TEMPLATE, active_tab='dashboard', emp_count=emp_count, total_kirim=total_kirim, total_chiqim=total_chiqim, sof_foyda=sof_foyda)

@app.route('/employees')
def employees():
    conn = get_db()
    emps = conn.execute("SELECT * FROM employees").fetchall()
    conn.close()
    return render_template_string(HTML_TEMPLATE, active_tab='employees', employees=emps)

@app.route('/employees/add', methods=['POST'])
def add_employee():
    name = request.form.get('name')
    position = request.form.get('position')
    phone = request.form.get('phone')
    if name:
        try:
            conn = get_db()
            conn.execute("INSERT INTO employees (name, position, phone, status) VALUES (?, ?, ?, 'ishda')", (name, position, phone))
            conn.commit()
            conn.close()
            return redirect(url_for('employees'))
        except Exception:
            pass
    return redirect(url_for('employees'))

@app.route('/salary')
def salary():
    conn = get_db()
    salaries = conn.execute("SELECT * FROM work_records ORDER BY work_date DESC").fetchall()
    conn.close()
    return render_template_string(HTML_TEMPLATE, active_tab='salary', salaries=salaries)

@app.route('/transactions')
def transactions():
    conn = get_db()
    txs = conn.execute("SELECT * FROM transactions ORDER BY tx_date DESC").fetchall()
    conn.close()
    return render_template_string(HTML_TEMPLATE, active_tab='transactions', transactions=txs, today=datetime.today().strftime('%Y-%m-%d'))

@app.route('/transactions/add', methods=['POST'])
def add_transaction():
    t_type = request.form.get('type')
    category = request.form.get('category')
    kontragent = request.form.get('kontragent')
    amount = request.form.get('amount')
    tx_date = request.form.get('tx_date')
    if category and amount and tx_date:
        try:
            conn = get_db()
            conn.execute("INSERT INTO transactions (type, category, kontragent, amount, tx_date) VALUES (?, ?, ?, ?, ?)",
                         (t_type, category, kontragent, float(amount), tx_date))
            conn.commit()
            conn.close()
        except Exception:
            pass
    return redirect(url_for('transactions'))

@app.route('/balance')
def balance():
    conn = get_db()
    savdo = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='kirim' AND (category LIKE '%savdo%' OR category LIKE '%Савдо%')").fetchone()[0] or 0
    tovar = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND (category LIKE '%tovar%' OR category LIKE '%товар%')").fetchone()[0] or 0
    ish_haqi = conn.execute("SELECT SUM(amount) FROM work_records").fetchone()[0] or 0
    kommunal = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND category LIKE '%kommunal%'").fetchone()[0] or 0
    arenda = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND category LIKE '%arenda%'").fetchone()[0] or 0
    marketing = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND category LIKE '%marketing%'").fetchone()[0] or 0
    remont = conn.execute("SELECT SUM(amount) FROM transactions WHERE type='chiqim' AND category LIKE '%remont%'").fetchone()[0] or 0
    nalog = savdo * 0.02
    yalpi_foyda = savdo - tovar
    jami_xarajat = ish_haqi + kommunal + arenda + marketing + remont + nalog
    sof_foyda = yalpi_foyda - jami_xarajat
    conn.close()
    
    bal = {
        'savdo': savdo,
        'tovar': tovar,
        'yalpi_foyda': yalpi_foyda,
        'ish_haqi': ish_haqi,
        'kommunal': kommunal,
        'arenda': arenda,
        'marketing': marketing,
        'remont': remont,
        'nalog': nalog,
        'jami_xarajat': jami_xarajat,
        'sof_foyda': sof_foyda
    }
    return render_template_string(HTML_TEMPLATE, active_tab='balance', balance=bal)

@app.route('/export')
def export_excel():
    wb = openpyxl.Workbook()
    conn = get_db()
    
    ws_emp = wb.active
    ws_emp.title = "Xodimlar"
    ws_emp.append(["ID", "Ism familiya", "Lavozimi", "Telefon", "Tug'ilgan kun", "Status"])
    for row in conn.execute("SELECT * FROM employees"):
        ws_emp.append([row["id"], row["name"], row["position"], row["phone"], row["birth_date"], row["status"]])
        
    ws_sal = wb.create_sheet(title="Ish haqi")
    ws_sal.append(["ID", "Xodim", "Sana", "Kirish", "Chiqish", "Kunbay summa", "Ishlagan vaqti"])
    for row in conn.execute("SELECT * FROM work_records"):
        ws_sal.append([row["id"], row["employee_name"], row["work_date"], row["check_in"], row["check_out"], row["amount"], row["worked_hours"]])
        
    ws_tx = wb.create_sheet(title="Kirim-Chiqim")
    ws_tx.append(["ID", "Turi", "Modda", "Kontragent", "Summa", "Sana"])
    for row in conn.execute("SELECT * FROM transactions"):
        ws_tx.append([row["id"], row["type"], row["category"], row["kontragent"], row["amount"], row["tx_date"]])
        
    conn.close()
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(output, as_attachment=True, download_name="Restoran_Hisobot.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == '__main__':
    print("Dastur ishga tushdi! Brauzerda oching: http://127.0.0.1:5000")
    app.run(debug=True, port=5000, use_reloader=False)
