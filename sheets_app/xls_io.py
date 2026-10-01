"""Excel (.xlsx) import / export."""
from datetime import date, datetime, time, timedelta

import openpyxl

from .storage import Workbook
from .formula import Evaluator, fmt


def _cell_text(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        if v.hour == v.minute == v.second == 0 and v.year > 1900:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, timedelta):
        total = int(v.total_seconds())
        h, m, s = total // 3600, (total % 3600) // 60, total % 60
        return f"{h}:{m:02d}:{s:02d}"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def import_excel(store: Workbook, path: str):
    """Excel faylni qiymatlar ko'rinishida import qiladi (formulalar emas).

    Nomlari mos varaq allaqachon mavjud bo'lsa — u tozalanib, qayta to'ldiriladi
    (takroriy varaq paydo bo'lmaydi).
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        title = ws.title[:30] or "Varaq"
        sid = store.sheet_id(title)
        if sid is None:
            sid = store.add_sheet(title)
        else:
            store.clear_sheet(sid)
        for row in ws.iter_rows():
            for cell in row:
                txt = _cell_text(cell.value)
                if txt is None or txt == "":
                    continue
                store.set_raw(sid, cell.row - 1, cell.column - 1, txt)
    wb.close()
    return len(wb.worksheets)


def export_excel(store: Workbook, path: str):
    """Barcha varaqni hisoblangan qiymatlar bilan Excel'ga yozadi."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ev = Evaluator(store)
    for sh in store.list_sheets():
        ws = wb.create_sheet(title=sh["name"][:31])
        r0, r1, c0, c1 = store.used_range(sh["id"])
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                raw = store.get_raw(sh["id"], r, c)
                if raw is None or raw == "":
                    continue
                v = ev.raw_value(sh["id"], r, c)
                if isinstance(v, str) and v in ("#CYCLE!", "#REF!", "#NAME?",
                                                "#VALUE!", "#DIV/0!", "#PARSE!"):
                    ws.cell(r + 1, c + 1, v)
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    ws.cell(r + 1, c + 1, v)
                elif isinstance(v, bool):
                    ws.cell(r + 1, c + 1, "TRUE" if v else "FALSE")
                else:
                    ws.cell(r + 1, c + 1, v)
    wb.save(path)
    wb.close()
