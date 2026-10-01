import openpyxl
import sqlite3
from datetime import datetime, time, timedelta
import os

EXCEL_PATH = r"C:\Users\HP Victus\Downloads\Ish haqi.xlsx"
DB_PATH = "restaurant.db"

def init_db(cursor):
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS employees (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE,
            position TEXT,
            phone TEXT,
            birth_date TEXT,
            status TEXT,
            daily_rate REAL DEFAULT 0
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS work_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            employee_name TEXT,
            work_date TEXT,
            check_in TEXT,
            check_out TEXT,
            amount REAL,
            worked_hours TEXT,
            UNIQUE(employee_name, work_date)
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            type TEXT,
            category TEXT,
            kontragent TEXT,
            amount REAL,
            tx_date TEXT
        )
    ''')

def parse_date(val):
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d')
    if isinstance(val, str):
        val = val.strip()
        for fmt in ('%m.%d.%Y', '%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y'):
            try:
                return datetime.strptime(val, fmt).strftime('%Y-%m-%d')
            except ValueError:
                continue
    return val

def run_import():
    if not os.path.exists(EXCEL_PATH):
        print(f"Xatolik: {EXCEL_PATH} topilmadi!")
        return

    wb_val = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    init_db(cursor)
    
    # Clear existing data to avoid duplicates when re-running
    cursor.execute('DELETE FROM employees')
    cursor.execute('DELETE FROM work_records')
    cursor.execute('DELETE FROM transactions')
    
    # 1. Kontragent (Xodimlar)
    if 'Kontragent' in wb_val.sheetnames:
        ws_k = wb_val['Kontragent']
        emp_count = 0
        for r in range(3, ws_k.max_row + 1):
            name = ws_k.cell(row=r, column=2).value
            if not name:
                continue
            name = str(name).strip()
            position = ws_k.cell(row=r, column=3).value
            phone = ws_k.cell(row=r, column=4).value
            birth = ws_k.cell(row=r, column=6).value
            status = ws_k.cell(row=r, column=7).value
            
            position = str(position).strip() if position else ""
            phone = str(phone).strip() if phone else ""
            birth_str = parse_date(birth) if birth else ""
            status_str = str(status).strip() if status else "ishda"
            
            cursor.execute('''
                INSERT OR IGNORE INTO employees (name, position, phone, birth_date, status)
                VALUES (?, ?, ?, ?, ?)
            ''', (name, position, phone, birth_str, status_str))
            emp_count += 1
        print(f"Kontragent (Xodimlar) import qilindi: {emp_count} ta.")

    # 2. Ish haqi
    total_ish_haqi = 0
    daily_sums = {}
    if 'ish haqi' in wb_val.sheetnames:
        ws_i = wb_val['ish haqi']
        date_cols = []
        for c in range(1, ws_i.max_column + 1):
            d_val = ws_i.cell(row=2, column=c).value
            if isinstance(d_val, datetime):
                date_cols.append((c, d_val.strftime('%Y-%m-%d')))
        
        work_records_count = 0
        for r in range(5, ws_i.max_row + 1):
            emp_name = ws_i.cell(row=r, column=2).value
            if not emp_name:
                continue
            emp_name = str(emp_name).strip()
            
            for col_idx, date_str in date_cols:
                check_in = ws_i.cell(row=r, column=col_idx).value
                check_out = ws_i.cell(row=r, column=col_idx+1).value
                amount = ws_i.cell(row=r, column=col_idx+2).value
                hours = ws_i.cell(row=r, column=col_idx+3).value
                
                if check_in is not None or check_out is not None or (amount is not None and amount > 0):
                    amt = float(amount) if amount and isinstance(amount, (int, float)) else 0.0
                    
                    ci_str = str(check_in) if check_in else ""
                    co_str = str(check_out) if check_out else ""
                    
                    if isinstance(check_in, time) and isinstance(check_out, time):
                        dt1 = datetime.combine(datetime.today(), check_in)
                        dt2 = datetime.combine(datetime.today(), check_out)
                        if dt2 < dt1:
                            dt2 += timedelta(days=1)
                        diff = dt2 - dt1
                        hours_str = str(diff)
                    else:
                        hours_str = str(hours) if hours else ""

                    if amt > 0 or ci_str:
                        total_ish_haqi += amt
                        daily_sums[date_str] = daily_sums.get(date_str, 0) + amt
                        
                        try:
                            cursor.execute('''
                                INSERT OR REPLACE INTO work_records (employee_name, work_date, check_in, check_out, amount, worked_hours)
                                VALUES (?, ?, ?, ?, ?, ?)
                            ''', (emp_name, date_str, ci_str, co_str, amt, hours_str))
                            work_records_count += 1
                        except Exception:
                            pass
        print(f"Ish haqi yozuvlari import qilindi: {work_records_count} ta. Jami ish haqi: {total_ish_haqi:,.0f} so'm")
        print("Kunlik ish haqi taqsimoti:", daily_sums)

    # 3. Kirim
    if 'Кирим' in wb_val.sheetnames:
        ws_kir = wb_val['Кирим']
        kirim_count = 0
        total_kirim = 0
        for r in range(2, ws_kir.max_row + 1):
            category = ws_kir.cell(row=r, column=4).value
            amount = ws_kir.cell(row=r, column=5).value
            tx_date = ws_kir.cell(row=r, column=6).value
            if category or amount:
                amt = float(amount) if amount and isinstance(amount, (int, float)) else 0.0
                d_str = parse_date(tx_date) if tx_date else ""
                cat_str = str(category).strip() if category else ""
                if amt > 0:
                    total_kirim += amt
                    cursor.execute('''
                        INSERT INTO transactions (type, category, kontragent, amount, tx_date)
                        VALUES (?, ?, ?, ?, ?)
                    ''', ('kirim', cat_str, '', amt, d_str))
                    kirim_count += 1
        print(f"Kirim import qilindi: {kirim_count} ta, Jami kirim: {total_kirim:,.0f} so'm")

    # 4. Chiqim
    if 'Чиким' in wb_val.sheetnames:
        ws_chi = wb_val['Чиким']
        chiqim_count = 0
        total_chiqim = 0
        for r in range(2, ws_chi.max_row + 1):
            kontr = ws_chi.cell(row=r, column=4).value
            category = ws_chi.cell(row=r, column=5).value
            amount = ws_chi.cell(row=r, column=6).value
            tx_date = ws_chi.cell(row=r, column=7).value
            if kontr or category or amount:
                amt = float(amount) if amount and isinstance(amount, (int, float)) else 0.0
                d_str = parse_date(tx_date) if tx_date else ""
                kontr_str = str(kontr).strip() if kontr else ""
                cat_str = str(category).strip() if category else ""
                if amt > 0:
                    total_chiqim += amt
                    cursor.execute('''
                        INSERT INTO transactions (type, category, kontragent, amount, tx_date)
                        VALUES (?, ?, ?, ?, ?)
                    ''', ('chiqim', cat_str, kontr_str, amt, d_str))
                    chiqim_count += 1
        print(f"Chiqim import qilindi: {chiqim_count} ta, Jami chiqim: {total_chiqim:,.0f} so'm")

    conn.commit()
    conn.close()
    print("Import muvaffaqiyatli yakunlandi!")

if __name__ == '__main__':
    run_import()
